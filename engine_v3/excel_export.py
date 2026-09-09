# =============================================================================
# engine_v3/excel_export.py — Export Excel (robust)
# - แสดงทุกคอลัมน์ (ซ่อนแค่ helper ล้วน)
# - จัด 3 โซน: Load Confirm → AP → AR
# - สี header: เทา / ฟ้า / เขียว (ตัวอักษรขาว)
# - auto-width + freeze header
# =============================================================================

from io import BytesIO
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


COLOR_LC = "374151"   # เทาเข้ม
COLOR_AP = "2563EB"   # ฟ้า
COLOR_AR = "059669"   # เขียว
WHITE    = "FFFFFF"

HIDDEN = {"_is_flat", "_pickup_str", "_ap_row", "_ar_row"}

# key ภายใน — ยังอยู่ในไฟล์ แต่ set column hidden (Unhide แล้วเจอ)
HIDE_IN_EXCEL = {
    "Pri-AP", "Mandate Key", "Sup-Carrier", "Sup-Truck", "AP Stop",
    "Final key", "Pri-AP2", "Final key2",
    "AR-Pri", "AR Stop", "AR-Final key", "AR-Pri2", "AR-Final key2",
    "AP Rate Type",   # ซ้ำกับ AP Rate Source (Source ละเอียดกว่า)
}

AR_HINTS = ["AR-Pri", "AR Stop", "AR-Final", "AR Prime", "AR Rate"]
AP_HINTS = ["Pri-AP", "Mandate Key", "Sup-Carrier", "Sup-Truck", "AP Stop",
            "Final key", "Prime Status", "Child Status", "Mandatory Status",
            "Carrier Status", "Truck Status", "AP Rate", "AP Effective", "AP Expiration"]
# ลำดับที่ต้องการภายในโซน AP (ตัวที่ไม่อยู่ในนี้ต่อท้ายตามเดิม)
AP_ORDER = [
    "Pri-AP", "Mandate Key", "Sup-Carrier", "Sup-Truck", "AP Stop",
    "Final key", "Pri-AP2", "Final key2",
    "Prime Status", "Child Status", "Mandatory Status", "Carrier Status", "Truck Status",
    "AP Effective Date", "AP Expiration Date",   # ← date ของตัวที่เลือก
    "AP Rate Charge", "AP Stop Charge",           # ← Stop Charge ต่อจาก Rate Charge
    "AP Rate Source", "AP Rate Type",
    "AP Rate From",   # ← ท้ายสุดโซน AP
]
# ลำดับภายในโซน AR
AR_ORDER = [
    "AR-Pri", "AR Stop", "AR-Final key", "AR-Pri2", "AR-Final key2",
    "AR Prime Status",
    "AR Rate Charge", "AR Stop Charge",           # ← Stop Charge ต่อจาก Rate Charge
    "AR Rate From",   # ← ท้ายสุดโซน AR
]


def _zone(col: str, orig_set: set) -> str:
    if col in orig_set:
        return "LC"
    for h in AR_HINTS:
        if h in col:
            return "AR"
    for h in AP_HINTS:
        if h in col:
            return "AP"
    return "LC"


def _order_zone(cols, preferred):
    """เรียง cols ตาม preferred ก่อน แล้วตัวที่เหลือต่อท้าย"""
    in_pref = [c for c in preferred if c in cols]
    rest = [c for c in cols if c not in preferred]
    return in_pref + rest


def _arrange(df, orig_cols):
    orig_cols = list(orig_cols) if orig_cols is not None else []
    orig_set = set(orig_cols)
    cols = [c for c in df.columns if c not in HIDDEN]
    lc, ap, ar = [], [], []
    for c in cols:
        z = _zone(str(c), orig_set)
        (lc if z == "LC" else ap if z == "AP" else ar).append(c)
    # LC เรียงตามต้นฉบับก่อน แล้ว Charge Type ต่อท้าย (ก่อนโซน AP)
    lc_orig = [c for c in orig_cols if c in lc]
    lc_extra = [c for c in lc if c not in orig_set and c != "Charge Type"]
    lc_ordered = lc_orig + lc_extra + (["Charge Type"] if "Charge Type" in lc else [])
    # AP / AR เรียงตามลำดับที่กำหนด (Rate From ท้ายสุด)
    ap_ordered = _order_zone(ap, AP_ORDER)
    ar_ordered = _order_zone(ar, AR_ORDER)
    order = lc_ordered + ap_ordered + ar_ordered
    order += [c for c in cols if c not in order]
    return df[order], orig_set


def to_styled_excel(df, original_cols=None, sheet_name="Result"):
    arranged, orig_set = _arrange(df, original_cols)

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        arranged.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]

        fills = {
            "LC": PatternFill(start_color=COLOR_LC, end_color=COLOR_LC, fill_type="solid"),
            "AP": PatternFill(start_color=COLOR_AP, end_color=COLOR_AP, fill_type="solid"),
            "AR": PatternFill(start_color=COLOR_AR, end_color=COLOR_AR, fill_type="solid"),
        }
        white_bold = Font(color=WHITE, bold=True)
        center = Alignment(horizontal="center", vertical="center")

        for idx, col in enumerate(arranged.columns, start=1):
            cell = ws.cell(row=1, column=idx)
            cell.fill = fills[_zone(str(col), orig_set)]
            cell.font = white_bold
            cell.alignment = center

        for idx, col in enumerate(arranged.columns, start=1):
            letter = get_column_letter(idx)
            max_len = len(str(col))
            try:
                for v in arranged[col].head(300):
                    if v is not None and not (isinstance(v, float) and pd.isna(v)):
                        L = len(str(v))
                        if L > max_len:
                            max_len = L
            except Exception:
                pass
            ws.column_dimensions[letter].width = min(max_len + 2, 50)

        # ── ซ่อน column key ภายใน (hidden=True — Unhide แล้วเจอ) ──
        for idx, col in enumerate(arranged.columns, start=1):
            if str(col) in HIDE_IN_EXCEL:
                ws.column_dimensions[get_column_letter(idx)].hidden = True

        ws.freeze_panes = "A2"

    return buf.getvalue()